import uuid
import gradio as gr
import torch
import csv
import openpyxl
from importlib import resources
from art import text2art
from openpyxl.utils.exceptions import *
from tapp.utils import TAPPModelInfer, TAPPException
from pathlib import Path
from tempfile import NamedTemporaryFile


tapp_model_infer: TAPPModelInfer


def calculate_physical_properties(*args):
    global tapp_model_infer
    input_field_names = ["Ti", "Al", "Cr", "Cu", "Fe", "Mo", "Ni", "Nb", "Si", "Sn", "V", "Zr", "N", "O", "C", "H", "B",
                         "Heat Treatment Temperature"]
    model_names = ["Thermal Expansion", "Density", "Thermal Conductivity", "Electrical Conductivity",
                   "Youngs Modulus", "Bulk Modulus", "Shear Modulus", "Poisson Ratio", "Specific Enthalpy",
                   "Specific Heat Capacity"]
    input_data = dict(zip(input_field_names, args))
    if all(input_data[element_name] < 1e-6 for element_name in input_field_names[1:12]):
        gr.Warning("请至少输入一个主要元素的含量。")
        return [None] * len(model_names)
    prop_values = []
    for model_name in model_names:
        try:
            prop_value = tapp_model_infer(model_name, input_data)
            prop_values.append(prop_value)
        except TAPPException as e:
            print(e)
            prop_values.append(None)
    return prop_values


def calculate_mechanical_properties(*args):
    global tapp_model_infer
    input_field_names = ["Ti", "Al", "Cr", "Cu", "Fe", "Mo", "Ni", "Nb", "Si", "Sn", "V", "Zr", "N", "O", "C", "H", "B",
                         "Heat Treatment Temperature", "Grain Size"]
    model_names = ["Yield Stress", "Tensile Stress", "Hardness", "Hall Petch"]
    input_data = dict(zip(input_field_names, args))
    if all(input_data[element_name] < 1e-6 for element_name in input_field_names[1:12]):
        gr.Warning("请至少输入一个主要元素的含量。")
        return [None] * len(model_names)
    prop_values = []
    for model_name in model_names:
        try:
            prop_value = tapp_model_infer(model_name, input_data)
            prop_values.append(prop_value)
        except TAPPException as e:
            print(e)
            prop_values.append(None)
    return prop_values


def read_csv_file(input_file_path: Path) -> list:
    input_data = []
    field_names = ["Ti", "Al", "Cr", "Cu", "Fe", "Mo", "Ni", "Nb", "Si", "Sn", "V", "Zr", "N", "O", "C", "H", "B",
                   "Heat Treatment Temperature", "Grain Size"]
    default_values = [100.0, 0.0, 0.0, 0.0, 0.0, 0.0, 0.0, 0.0, 0.0, 0.0, 0.0, 0.0, 0.0, 0.0, 0.0, 0.0, 0.0, 600.0, 10.0]
    with open(input_file_path, "r", newline="", encoding="utf-8") as csv_file:
        dict_reader = csv.DictReader(csv_file)
        for row in dict_reader:
            dict_data = {}
            for field_name, default_value in zip(field_names, default_values):
                if field_name in row:
                    if row[field_name] != "":
                        dict_data[field_name] = float(row[field_name])
                    else:
                        dict_data[field_name] = default_value
                else:
                    dict_data[field_name] = default_value
            input_data.append(dict_data)
    return input_data


def read_xlsx_file(input_file_path: Path) -> list:
    input_data = []
    field_names = ["Ti", "Al", "Cr", "Cu", "Fe", "Mo", "Ni", "Nb", "Si", "Sn", "V", "Zr", "N", "O", "C", "H", "B",
                   "Heat Treatment Temperature", "Grain Size"]
    default_values = [100.0, 0.0, 0.0, 0.0, 0.0, 0.0, 0.0, 0.0, 0.0, 0.0, 0.0, 0.0, 0.0, 0.0, 0.0, 0.0, 0.0, 600.0,
                      10.0]
    workbook = openpyxl.load_workbook(input_file_path, read_only=True)
    worksheet = workbook.active
    header = None
    for index, row in enumerate(worksheet.iter_rows(values_only=True)):
        if index == 0:
            header = [str(item) if item else str(uuid.uuid4()) for item in row]
        else:
            row = dict(zip(header, row))
            dict_data = {}
            for field_name, default_value in zip(field_names, default_values):
                if field_name in row:
                    if row[field_name] is not None:
                        dict_data[field_name] = float(row[field_name])
                    else:
                        dict_data[field_name] = default_value
                else:
                    dict_data[field_name] = default_value
            input_data.append(dict_data)
    return input_data


def write_csv_file(input_file_path: Path, out_data: dict) -> str:
    with NamedTemporaryFile(delete=False, prefix=f"{input_file_path.stem}_", suffix=".csv",
                            mode="w", newline="", encoding="utf-8") as output_file:
        with open(input_file_path, "r", newline="", encoding="utf-8") as input_file:
            dict_reader = csv.DictReader(input_file)
            field_names = list(dict_reader.fieldnames) + list(out_data.keys())
            dict_writer = csv.DictWriter(output_file, field_names)
            dict_writer.writeheader()
            for index, row in enumerate(dict_reader):
                for k, v in out_data.items():
                    row[k] = f"{v[index]}"
                dict_writer.writerow(row)
        return output_file.name


def write_xlsx_file(input_file_path: Path, out_data: dict) -> str:
    output_file = NamedTemporaryFile(delete=False, prefix=f"{input_file_path.stem}_", suffix=".xlsx", mode="w")
    output_file_path = output_file.name
    output_file.close()
    input_workbook = openpyxl.load_workbook(input_file_path, read_only=True)
    input_worksheet = input_workbook.active
    output_workbook = openpyxl.Workbook()
    output_worksheet = output_workbook.active
    for row in input_worksheet.iter_rows(values_only=True):
        output_worksheet.append(row)
    col_index = output_worksheet.max_column
    for property_name, property_values in out_data.items():
        col_index += 1
        output_worksheet.cell(row=1, column=col_index, value=property_name)
        for row_index, property_value in enumerate(property_values, start=2):
            output_worksheet.cell(row=row_index, column=col_index, value=f"{property_value}")
    output_workbook.save(output_file_path)
    return output_file_path


def batch_process(*args):
    global tapp_model_infer
    property_names = ["Thermal Expansion", "Density", "Thermal Conductivity", "Electrical Conductivity",
                      "Youngs Modulus", "Bulk Modulus", "Shear Modulus", "Poisson Ratio", "Specific Enthalpy",
                      "Specific Heat Capacity", "Yield Stress", "Tensile Stress", "Hardness", "Hall Petch"]
    property_names_zh = ["热膨胀系数", "密度", "热导率", "电导率", "杨氏模量", "体积模量", "剪切模量", "泊松比", "比焓",
                         "比热容", "屈服强度", "抗拉强度", "硬度", "霍尔佩奇系数"]
    selected_properties = [property_names[property_names_zh.index(property_name)] for property_name in args[0]]
    input_file_paths = args[1]
    if len(selected_properties) == 0:
        gr.Warning("请至少选择一个性能")
        return None
    if not input_file_paths:
        gr.Warning("请上传至少一个 CSV 或 XLSX 文件")
        return None
    output_files = []
    for input_file_path in input_file_paths:
        input_file_path = Path(input_file_path)
        if input_file_path.match("*.csv"):
            try:
                input_data = read_csv_file(input_file_path)
                if len(input_data) == 0:
                    gr.Warning(f"空的 CSV 文件：\"{input_file_path.name}\"")
                    continue
                output_data = {}
                for property_name in selected_properties:
                    try:
                        property_values = tapp_model_infer(property_name, input_data)
                        output_data[property_name] = property_values
                    except TAPPException as e:
                        print(e)
                        continue
                output_file_path = write_csv_file(input_file_path, output_data)
                output_files.append(output_file_path)
            except csv.Error as err:
                gr.Warning(f"错误的 CSV 文件：\"{input_file_path.name}\"")
                print(err)
            except Exception as err:
                gr.Warning(f"未知错误")
                print(err)
        elif input_file_path.match("*.xlsx"):
            try:
                input_data = read_xlsx_file(input_file_path)
                if len(input_data) == 0:
                    gr.Warning(f"空的 XLSX 文件：\"{input_file_path.name}\"")
                    continue
                output_data = {}
                for property_name in selected_properties:
                    try:
                        property_values = tapp_model_infer(property_name, input_data)
                        output_data[property_name] = property_values
                    except TAPPException as e:
                        print(e)
                        continue
                output_file_path = write_xlsx_file(input_file_path, output_data)
                output_files.append(output_file_path)
            except (CellCoordinatesException, IllegalCharacterError, NamedRangeException, SheetTitleException,
                    InvalidFileException, ReadOnlyWorkbookException, WorkbookAlreadySaved) as err:
                gr.Error(f"错误的 XLSX 文件：\"{input_file_path.name}\"")
                print(err)
            except Exception as err:
                gr.Error(f"未知错误")
                print(err)
    return output_files


def init_models():
    global tapp_model_infer
    device = torch.device("cuda" if torch.cuda.is_available() else "cpu")
    property_names = ["Thermal Expansion", "Density", "Thermal Conductivity", "Electrical Conductivity",
                      "Youngs Modulus", "Bulk Modulus", "Shear Modulus", "Poisson Ratio", "Specific Enthalpy",
                      "Specific Heat Capacity", "Yield Stress", "Tensile Stress", "Hardness", "Hall Petch"]
    model_files = {
        property_name: str(resources.files("tapp.data").joinpath(f"MoE2({property_name}).pth"))
        for property_name in property_names
    }
    norm_files = {
        property_name: str(resources.files("tapp.data").joinpath(f"norm_param({property_name}).json"))
        for property_name in property_names
    }
    tapp_model_infer = TAPPModelInfer(model_files=model_files, norm_files=norm_files, device=device)


def run_gradio():
    print(text2art("TAPP"))
    init_models()
    with (gr.Blocks(title="钛合金性能预测模块") as index):
        gr.Markdown("## 钛合金性能预测模块")
        with gr.Tabs():
            with gr.Tab("物理性能预测"):
                with gr.Row():
                    with gr.Column():
                        gr.Markdown("### 元素组成")
                        with gr.Row():
                            elements_p = [
                                gr.Number(
                                    label=f"{element_name} (wt%)",
                                    value=default_value,
                                    minimum=min_value,
                                    maximum=max_value,
                                    interactive=interactive,
                                    precision=2
                                )
                                for element_name, default_value, min_value, max_value, interactive in [
                                    ("Ti", 100.0, 25.0, 100.0, False),
                                    ("Al", 0.0, 0.0, 8.5, True),
                                    ("Cr", 0.0, 0.0, 12.0, True),
                                    ("Cu", 0.0, 0.0, 4.0, True),
                                    ("Fe", 0.0, 0.0, 12.0, True),
                                    ("Mo", 0.0, 0.0, 17.0, True),
                                    ("Ni", 0.0, 0.0, 12.0, True),
                                    ("Nb", 0.0, 0.0, 12.0, True),
                                    ("Si", 0.0, 0.0, 3.0, True),
                                    ("Sn", 0.0, 0.0, 12.0, True),
                                    ("V", 0.0, 0.0, 17.0, True),
                                    ("Zr", 0.0, 0.0, 12.0, True),
                                    ("N", 0.0, 0.0, 0.5, True),
                                    ("O", 0.0, 0.0, 0.5, True),
                                    ("C", 0.0, 0.0, 0.4, True),
                                    ("H", 0.0, 0.0, 0.2, True),
                                    ("B", 0.0, 0.0, 0.4, True)
                                ]
                            ]
                        gr.Markdown("### 工艺参数")
                        heat_treatment_temperature_p = gr.Number(
                            label="热处理温度 (°C)", value=600, minimum=25, maximum=1800, interactive=True, precision=2
                        )
                        with gr.Row():
                            clear_btn_p = gr.Button("重置", interactive=True)
                            calculate_btn_p = gr.Button("提交", interactive=True)
                    with gr.Column():
                        gr.Markdown("### 物理性能")
                        properties_p = [
                            gr.Number(label=f"{property_name} ({property_unit})" if property_unit else property_name,
                                      interactive=False)
                            for property_name, property_unit in [
                                ("热膨胀系数", "K⁻¹"),
                                ("密度", "g/cm³"),
                                ("热导率", "W/m·K"),
                                ("电导率", "S/m"),
                                ("杨氏模量", "GPa"),
                                ("体积模量", "GPa"),
                                ("剪切模量", "GPa"),
                                ("泊松比", None),
                                ("比焓", "J/g"),
                                ("比热容", "J/g·K"),
                            ]
                        ]
            with gr.Tab("力学性能预测"):
                with gr.Row():
                    with gr.Column():
                        gr.Markdown("### 元素组成")
                        with gr.Row():
                            elements_m = [
                                gr.Number(
                                    label=f"{element_name} (wt%)",
                                    value=default_value,
                                    minimum=min_value,
                                    maximum=max_value,
                                    interactive=interactive,
                                    precision=2
                                )
                                for element_name, default_value, min_value, max_value, interactive in [
                                    ("Ti", 100.0, 25.0, 100.0, False),
                                    ("Al", 0.0, 0.0, 8.5, True),
                                    ("Cr", 0.0, 0.0, 12.0, True),
                                    ("Cu", 0.0, 0.0, 4.0, True),
                                    ("Fe", 0.0, 0.0, 12.0, True),
                                    ("Mo", 0.0, 0.0, 17.0, True),
                                    ("Ni", 0.0, 0.0, 12.0, True),
                                    ("Nb", 0.0, 0.0, 12.0, True),
                                    ("Si", 0.0, 0.0, 3.0, True),
                                    ("Sn", 0.0, 0.0, 12.0, True),
                                    ("V", 0.0, 0.0, 17.0, True),
                                    ("Zr", 0.0, 0.0, 12.0, True),
                                    ("N", 0.0, 0.0, 0.5, True),
                                    ("O", 0.0, 0.0, 0.5, True),
                                    ("C", 0.0, 0.0, 0.4, True),
                                    ("H", 0.0, 0.0, 0.2, True),
                                    ("B", 0.0, 0.0, 0.4, True)
                                ]
                            ]
                        gr.Markdown("### 工艺参数")
                        heat_treatment_temperature_m = gr.Number(
                            label="热处理温度 (°C)", value=600, minimum=25, maximum=1800, interactive=True, precision=2
                        )
                        gr.Markdown("### 晶粒尺寸")
                        grain_size_m = gr.Number(label="Alpha & Beta (μm)", value=10, minimum=0, maximum=200,
                                                 interactive=True, precision=2)
                        with gr.Row():
                            clear_btn_m = gr.Button("重置", interactive=True)
                            calculate_btn_m = gr.Button("提交", interactive=True)
                    with gr.Column():
                        gr.Markdown("### 力学性能")
                        properties_m = [
                            gr.Number(label=f"{property_name} ({property_unit})", interactive=False)
                            for property_name, property_unit in [
                                ("屈服强度", "MPa"),
                                ("抗拉强度", "MPa"),
                                ("硬度", "VPN"),
                                ("霍尔佩奇系数", "MPa·√m")
                            ]
                        ]
            with gr.Tab("批量模式"):
                with gr.Row():
                    with gr.Column():
                        up_files = gr.File(label="上传 CSV/XLSX 文件", file_types=[".csv", ".xlsx"],
                                             file_count="multiple", type="filepath", interactive=True)
                        down_files = gr.File(label="下载 CSV/XLSX 文件", file_types=[".csv", ".xlsx"],
                                               file_count="multiple", type="filepath", interactive=False)
                    with gr.Column():
                        checkboxes = gr.CheckboxGroup(choices=["热膨胀系数", "密度", "热导率", "电导率", "杨氏模量", "体积模量",
                                                               "剪切模量", "泊松比", "比焓", "比热容", "屈服强度", "抗拉强度",
                                                               "硬度", "霍尔佩奇系数"],
                                                      label="选择性能",
                                                      interactive=True)
                        with gr.Row():
                            clean_btn_b = gr.Button("重置", interactive=True)
                            batch_process_btn = gr.Button("提交", interactive=True)
                            select_all_btn = gr.Button("全选", interactive=True)
            with gr.Tab("关于"):
                with gr.Row():
                    with gr.Column():
                        gr.Image(
                            value=str(resources.files("tapp.data").joinpath("logo.png")),
                            width=200,
                            height=200,
                            show_label=False,
                            interactive=False
                        )
                        gr.Markdown("**TAPP（Titanium Alloy Property Predictor）** 是一款利用机器学习技术进行钛合金材料属性预测的应用程序。")
                    gr.File(
                        value=[
                            str(resources.files("tapp.data").joinpath("tapp_template.csv")),
                            str(resources.files("tapp.data").joinpath("tapp_template.xlsx"))
                        ],
                        label="下载批处理模板文件",
                        file_count="multiple",
                        type="filepath",
                        interactive=False
                    )
        for element in elements_p[1:]:
            element.change(
                fn=lambda *args: 100.0 - sum(args),
                inputs=elements_p[1:],
                outputs=elements_p[0]
            )
        clear_btn_p.click(
            fn=lambda: [100, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 600, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0],
            outputs=elements_p + [heat_treatment_temperature_p] + properties_p
        )
        calculate_btn_p.click(
            fn=calculate_physical_properties,
            inputs=elements_p + [heat_treatment_temperature_p],
            outputs=properties_p
        )
        for element in elements_m[1:]:
            element.change(
                fn=lambda *args: 100 - sum(args),
                inputs=elements_m[1:],
                outputs=elements_m[0]
            )
        clear_btn_m.click(
            fn=lambda: [100, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 600, 10, 0, 0, 0, 0],
            outputs=elements_m + [heat_treatment_temperature_m, grain_size_m] + properties_m
        )
        calculate_btn_m.click(
            fn=calculate_mechanical_properties,
            inputs=elements_m + [heat_treatment_temperature_m, grain_size_m],
            outputs=properties_m
        )
        batch_process_btn.click(
            fn=batch_process,
            inputs=[checkboxes, up_files],
            outputs=down_files
        )
        select_all_btn.click(
            fn=lambda selected_properties: list(
                {"热膨胀系数", "密度", "热导率", "电导率", "杨氏模量", "体积模量", "剪切模量", "泊松比", "比焓",
                 "比热容", "屈服强度", "抗拉强度", "硬度", "霍尔佩奇系数"} - set(selected_properties)
            ),
            inputs=checkboxes,
            outputs=checkboxes
        )
        clean_btn_b.click(
            fn=lambda: [None, None, None],
            outputs=[up_files, down_files, checkboxes]
        )
        index.launch()
