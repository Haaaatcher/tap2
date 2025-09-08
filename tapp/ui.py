import gradio as gr
import csv
import openpyxl
from importlib import resources
from art import text2art
from gradio.utils import NamedString
from tapp.core import TAPPModelInfer
from pathlib import Path
from tempfile import NamedTemporaryFile


infer: TAPPModelInfer


def get_phys_prop_values(Ti: int|float|None, Al: int|float|None, Cr: int|float|None, Cu: int|float|None,
        Fe: int|float|None, Mo: int|float|None, Ni: int|float|None, Nb: int|float|None, Si: int|float|None,
        Sn: int|float|None, V: int|float|None, Zr: int|float|None, N: int|float|None, O: int|float|None,
        C: int|float|None, H: int|float|None, B: int|float|None, Mn: int|float|None, Ta: int|float|None,
        Bi: int|float|None, Co: int|float|None, htt: int|float|None) -> list[float]:
    """
    调用TAPP获取物理性能预测结果。
    :param Ti: 钛的质量分数
    :param Al: 铝的质量分数
    :param Cr: 铬的质量分数
    :param Cu: 铜的质量分数
    :param Fe: 铁的质量分数
    :param Mo: 钼的质量分数
    :param Ni: 镍的质量分数
    :param Nb: 铌的质量分数
    :param Si: 硅的质量分数
    :param Sn: 锡的质量分数
    :param V: 钒的质量分数
    :param Zr: 锆的质量分数
    :param N: 氮的质量分数
    :param O: 氧的质量分数
    :param C: 碳的质量分数
    :param H: 氢的质量分数
    :param B: 硼的质量分数
    :param Mn: 锰的质量分数
    :param Ta: 钽的质量分数
    :param Bi: 铋的质量分数
    :param Co: 钴的质量分数
    :param htt: 热处理温度
    :return: 物理性能列表
    """
    global infer
    # 构造输入字典
    input_ = {}
    if Ti is not None:
        input_["Ti"] = Ti
    if Al is not None:
        input_["Al"] = Al
    if Cr is not None:
        input_["Cr"] = Cr
    if Cu is not None:
        input_["Cu"] = Cu
    if Fe is not None:
        input_["Fe"] = Fe
    if Mo is not None:
        input_["Mo"] = Mo
    if Ni is not None:
        input_["Ni"] = Ni
    if Nb is not None:
        input_["Nb"] = Nb
    if Si is not None:
        input_["Si"] = Si
    if Sn is not None:
        input_["Sn"] = Sn
    if V is not None:
        input_["V"] = V
    if Zr is not None:
        input_["Zr"] = Zr
    if N is not None:
        input_["N"] = N
    if O is not None:
        input_["O"] = O
    if C is not None:
        input_["C"] = C
    if H is not None:
        input_["H"] = H
    if B is not None:
        input_["B"] = B
    if Mn is not None:
        input_["Mn"] = Mn
    if Ta is not None:
        input_["Ta"] = Ta
    if Bi is not None:
        input_["Bi"] = Bi
    if Co is not None:
        input_["Co"] = Co
    if htt is not None:
        input_["Heat Treatment Temperature"] = htt
    phys_prop_values = []
    for prop_name in ["Thermal Expansion", "Density", "Thermal Conductivity", "Electrical Conductivity",
                      "Youngs Modulus", "Bulk Modulus", "Shear Modulus", "Poisson Ratio", "Specific Enthalpy",
                      "Specific Heat Capacity"]:
        prop_value = infer(prop_name, input_)
        phys_prop_values.append(prop_value)
    # 修正单位
    phys_prop_values[0] *= 1e6
    return phys_prop_values


def get_mech_prop_values(Ti: int|float|None, Al: int|float|None, Cr: int|float|None, Cu: int|float|None,
        Fe: int|float|None, Mo: int|float|None, Ni: int|float|None, Nb: int|float|None, Si: int|float|None,
        Sn: int|float|None, V: int|float|None, Zr: int|float|None, N: int|float|None, O: int|float|None,
        C: int|float|None, H: int|float|None, B: int|float|None, Mn: int|float|None, Ta: int|float|None,
        Bi: int|float|None, Co: int|float|None, htt: int|float|None) -> list[float]:
    """
    调用TAPP获取力学性能预测结果。
    :param Ti: 钛的质量分数
    :param Al: 铝的质量分数
    :param Cr: 铬的质量分数
    :param Cu: 铜的质量分数
    :param Fe: 铁的质量分数
    :param Mo: 钼的质量分数
    :param Ni: 镍的质量分数
    :param Nb: 铌的质量分数
    :param Si: 硅的质量分数
    :param Sn: 锡的质量分数
    :param V: 钒的质量分数
    :param Zr: 锆的质量分数
    :param N: 氮的质量分数
    :param O: 氧的质量分数
    :param C: 碳的质量分数
    :param H: 氢的质量分数
    :param B: 硼的质量分数
    :param Mn: 锰的质量分数
    :param Ta: 钽的质量分数
    :param Bi: 铋的质量分数
    :param Co: 钴的质量分数
    :param htt: 热处理温度
    :return: 力学性能
    """
    global infer
    # 构造输入字典
    input_ = {}
    if Ti is not None:
        input_["Ti"] = Ti
    if Al is not None:
        input_["Al"] = Al
    if Cr is not None:
        input_["Cr"] = Cr
    if Cu is not None:
        input_["Cu"] = Cu
    if Fe is not None:
        input_["Fe"] = Fe
    if Mo is not None:
        input_["Mo"] = Mo
    if Ni is not None:
        input_["Ni"] = Ni
    if Nb is not None:
        input_["Nb"] = Nb
    if Si is not None:
        input_["Si"] = Si
    if Sn is not None:
        input_["Sn"] = Sn
    if V is not None:
        input_["V"] = V
    if Zr is not None:
        input_["Zr"] = Zr
    if N is not None:
        input_["N"] = N
    if O is not None:
        input_["O"] = O
    if C is not None:
        input_["C"] = C
    if H is not None:
        input_["H"] = H
    if B is not None:
        input_["B"] = B
    if Mn is not None:
        input_["Mn"] = Mn
    if Ta is not None:
        input_["Ta"] = Ta
    if Bi is not None:
        input_["Bi"] = Bi
    if Co is not None:
        input_["Co"] = Co
    if htt is not None:
        input_["Heat Treatment Temperature"] = htt
    mech_prop_values = []
    for prop_name in ["Yield Stress", "Tensile Stress", "Hardness", "Hall Petch"]:
        prop_value = infer(prop_name, input_)
        mech_prop_values.append(prop_value)
    return mech_prop_values


def read_csv(file_path: Path) -> list[dict[str, int | float]]:
    """
    读取CSV文件，返回TAPP输入列表。
    :param file_path: CSV文件路径
    :return: TAPP输入列表
    """
    field_names = ["Ti", "Al", "Cr", "Cu", "Fe", "Mo", "Ni", "Nb", "Si", "Sn", "V", "Zr", "N",
                   "O", "C", "H", "B", "Mn", "Ta", "Bi", "Co", "Heat Treatment Temperature"]
    inputs = []
    with open(file_path, "r", newline="", encoding="utf-8") as csv_file:
        reader = csv.reader(csv_file)
        col_names = next(reader)
        for row in reader:
            input_ = {}
            for col_name, cell_value in zip(col_names, row):
                if col_name in field_names and cell_value != "":
                    input_[col_name] = float(cell_value)
            inputs.append(input_)
    return inputs


def read_xlsx(file_path: Path) -> list[dict[str, int | float]]:
    """
    读取XLSX文件，返回TAPP输入列表。
    :param file_path: XLSX文件路径
    :return: TAPP输入列表
    """
    field_names = ["Ti", "Al", "Cr", "Cu", "Fe", "Mo", "Ni", "Nb", "Si", "Sn", "V", "Zr", "N",
                   "O", "C", "H", "B", "Mn", "Ta", "Bi", "Co", "Heat Treatment Temperature"]
    inputs = []
    wb = openpyxl.load_workbook(file_path, read_only=True)
    ws = wb.active
    table = list(ws.iter_rows(values_only=True))
    col_names = table[0]
    for row in table[1:]:
        input_ = {}
        for col_name, cell_value in zip(col_names, row):
            if col_name in field_names and cell_value is not None:
                input_[col_name] = float(cell_value)
        inputs.append(input_)
    return inputs


def write_csv(input_path: Path, outputs_dict: dict[str, list[float]]) -> Path:
    """
    将TAPP的预测结果写入输出CSV文件。
    :param input_path: 输入CSV文件路径
    :param outputs_dict: TAPP预测结果
    :return: 输出CSV文件路径
    """
    # 创建新的CSV文件
    with NamedTemporaryFile(delete=False,
                            prefix=f"{input_path.stem}_",
                            suffix=".csv",
                            mode="w",
                            newline="",
                            encoding="utf-8") as output_file:
        with open(input_path, "r", newline="", encoding="utf-8") as input_file:
            reader = csv.reader(input_file)
            writer = csv.writer(output_file)
            col_names = next(reader) + list(outputs_dict.keys())
            writer.writerow(col_names)
            row_num = 0
            for row in reader:
                for prop_name, outputs in outputs_dict.items():
                    prop_value = outputs[row_num]
                    # 修正单位
                    if prop_name == "Thermal Expansion":
                        prop_value = prop_value * 1e6
                    row.append(f"{prop_value :.3f}")
                writer.writerow(row)
                row_num += 1
        return Path(output_file.name)


def write_xlsx(input_path: Path, outputs_dict: dict[str, list[float]]) -> Path:
    """
    将TAPP的预测结果写入输出XLSX文件。
    :param input_path: 输入XLSX文件路径
    :param outputs_dict: TAPP预测结果
    :return: 输出XLSX文件路径
    """
    # 创建新的XLSX文件
    output_file = NamedTemporaryFile(delete=False, prefix=f"{input_path.stem}_", suffix=".xlsx", mode="w")
    output_path = Path(output_file.name)
    output_file.close()
    input_wb = openpyxl.load_workbook(input_path, read_only=True)
    input_ws = input_wb.active
    output_wb = openpyxl.Workbook()
    output_ws = output_wb.active
    # 写入原始数据
    for row in input_ws.iter_rows(values_only=True):
        output_ws.append(row)
    col_num = output_ws.max_column
    for prop_name, outputs in outputs_dict.items():
        col_num += 1
        output_ws.cell(row=1, column=col_num, value=prop_name)
        for row_num, prop_value in enumerate(outputs, start=2):
            # 修正单位
            if prop_name == "Thermal Expansion":
                prop_value = prop_value * 1e6
            output_ws.cell(row=row_num, column=col_num, value=f"{prop_value:.3f}")
    output_wb.save(output_path)
    return output_path


def batch_proc(prop_names: list[str], input_files: list[NamedString]) -> list[str]:
    """
    批量处理CSV或XLSX文件，调用TAPP的批处理功能进行属性预测，返回包含预测结果的文件路径列表。
    :param prop_names: 选择性能列表
    :param input_files: 输入文件路径列表
    :return: 输出文件路径列表
    """
    global infer
    prop_map = {
        "热膨胀系数": "Thermal Expansion",
        "密度": "Density",
        "热导率": "Thermal Conductivity",
        "电导率": "Electrical Conductivity",
        "杨氏模量": "Youngs Modulus",
        "体积模量": "Bulk Modulus",
        "剪切模量": "Shear Modulus",
        "泊松比": "Poisson Ratio",
        "比焓": "Specific Enthalpy",
        "比热容": "Specific Heat Capacity",
        "屈服强度": "Yield Stress",
        "抗拉强度": "Tensile Stress",
        "硬度": "Hardness",
        "霍尔佩奇系数": "Hall Petch"
    }
    if len(prop_names) == 0:
        gr.Warning("请至少选择一个性能")
        return []
    if len(input_files) == 0:
        gr.Warning("请上传至少一个 CSV 或 XLSX 文件")
        return []
    output_paths = []
    for input_path in input_files:
        input_path = Path(input_path)
        if input_path.match("*.csv"):
            reader = read_csv
            writer = write_csv
        elif input_path.match("*.xlsx"):
            reader = read_xlsx
            writer = write_xlsx
        else:
            gr.Warning(f"不支持的文件格式：\"{input_path.name}\"")
            continue
        inputs = reader(input_path)
        if not inputs:
            gr.Warning(f"空的 CSV 文件：\"{input_path.name}\"")
            continue
        outputs_dict = {}
        for prop_name in prop_names:
            outputs = infer(prop_map[prop_name], inputs)
            outputs_dict[prop_map[prop_name]] = outputs
        output_path = str(writer(input_path, outputs_dict))
        output_paths.append(output_path)
    return output_paths


def run_gradio():
    """
    基于Gradio的用户图形界面，支持物理性能预测、力学性能预测、批量模式、模板下载。
    :return: None
    """
    global infer
    print(text2art("TAPP"))
    infer = TAPPModelInfer()
    with gr.Blocks(title="钛合金性能预测模块") as index:
        # 定义界面布局
        gr.Markdown("## 钛合金性能预测模块")
        with gr.Tabs():
            with gr.Tab("物理性能"):
                with gr.Row():
                    with gr.Column():
                        gr.Markdown("### 元素组成")
                        with gr.Row():
                            phys_Ti = gr.Number(label=f"Ti (wt%)", value=100, minimum=0, maximum=100, interactive=True)
                            phys_Al = gr.Number(label=f"Al (wt%)", value=0, minimum=0, maximum=100, interactive=True)
                            phys_Cr = gr.Number(label=f"Cr (wt%)", value=0, minimum=0, maximum=100, interactive=True)
                            phys_Cu = gr.Number(label=f"Cu (wt%)", value=0, minimum=0, maximum=100, interactive=True)
                            phys_Fe = gr.Number(label=f"Fe (wt%)", value=0, minimum=0, maximum=100, interactive=True)
                            phys_Mo = gr.Number(label=f"Mo (wt%)", value=0, minimum=0, maximum=100, interactive=True)
                            phys_Ni = gr.Number(label=f"Ni (wt%)", value=0, minimum=0, maximum=100, interactive=True)
                            phys_Nb = gr.Number(label=f"Nb (wt%)", value=0, minimum=0, maximum=100, interactive=True)
                            phys_Si = gr.Number(label=f"Si (wt%)", value=0, minimum=0, maximum=100, interactive=True)
                            phys_Sn = gr.Number(label=f"Sn (wt%)", value=0, minimum=0, maximum=100, interactive=True)
                            phys_V = gr.Number(label=f"V (wt%)", value=0, minimum=0, maximum=100, interactive=True)
                            phys_Zr = gr.Number(label=f"Zr (wt%)", value=0, minimum=0, maximum=100, interactive=True)
                            phys_N = gr.Number(label=f"N (wt%)", value=0, minimum=0, maximum=100, interactive=True)
                            phys_O = gr.Number(label=f"O (wt%)", value=0, minimum=0, maximum=100, interactive=True)
                            phys_C = gr.Number(label=f"C (wt%)", value=0, minimum=0, maximum=100, interactive=True)
                            phys_H = gr.Number(label=f"H (wt%)", value=0, minimum=0, maximum=100, interactive=True)
                            phys_B = gr.Number(label=f"B (wt%)", value=0, minimum=0, maximum=100, interactive=True)
                            phys_Mn = gr.Number(label=f"Mn (wt%)", value=0, minimum=0, maximum=100, interactive=True)
                            phys_Ta = gr.Number(label=f"Ta (wt%)", value=0, minimum=0, maximum=100, interactive=True)
                            phys_Bi = gr.Number(label=f"Bi (wt%)", value=0, minimum=0, maximum=100, interactive=True)
                            phys_Co = gr.Number(label=f"Co (wt%)", value=0, minimum=0, maximum=100, interactive=True)
                        gr.Markdown("### 工艺参数")
                        phys_htt = gr.Number(label="热处理温度 (°C)", value=600, interactive=True)
                        with gr.Row():
                            phys_clear_btn = gr.Button("重置", interactive=True)
                            phys_run_btn = gr.Button("提交", interactive=True)
                    with gr.Column():
                        gr.Markdown("### 物理性能")
                        phys_te = gr.Number(label=r"热膨胀系数 (10^-6/K)", value=None, interactive=False, precision=3)
                        phys_d = gr.Number(label="密度 (g/cm^3)", value=None, interactive=False, precision=3)
                        phys_tc = gr.Number(label="热导率 (W/m·K)", value=None, interactive=False, precision=3)
                        phys_ec = gr.Number(label="电导率 (S/m)", value=None, interactive=False, precision=3)
                        phys_ym = gr.Number(label="杨氏模量 (GPa)", value=None, interactive=False, precision=3)
                        phys_bm = gr.Number(label="体积模量 (GPa)", value=None, interactive=False, precision=3)
                        phys_sm = gr.Number(label="剪切模量 (GPa)", value=None, interactive=False, precision=3)
                        phys_pr = gr.Number(label="泊松比", value=None, interactive=False, precision=3)
                        phys_se = gr.Number(label="比焓 (J/g)", value=None, interactive=False, precision=3)
                        phys_shc = gr.Number(label="比热容 (J/g·K)", value=None, interactive=False, precision=3)
            with gr.Tab("力学性能"):
                with gr.Row():
                    with gr.Column():
                        gr.Markdown("### 元素组成")
                        with gr.Row():
                            mech_Ti = gr.Number(label=f"Ti (wt%)", value=100, minimum=0, maximum=100, interactive=True)
                            mech_Al = gr.Number(label=f"Al (wt%)", value=0, minimum=0, maximum=100, interactive=True)
                            mech_Cr = gr.Number(label=f"Cr (wt%)", value=0, minimum=0, maximum=100, interactive=True)
                            mech_Cu = gr.Number(label=f"Cu (wt%)", value=0, minimum=0, maximum=100, interactive=True)
                            mech_Fe = gr.Number(label=f"Fe (wt%)", value=0, minimum=0, maximum=100, interactive=True)
                            mech_Mo = gr.Number(label=f"Mo (wt%)", value=0, minimum=0, maximum=100, interactive=True)
                            mech_Ni = gr.Number(label=f"Ni (wt%)", value=0, minimum=0, maximum=100, interactive=True)
                            mech_Nb = gr.Number(label=f"Nb (wt%)", value=0, minimum=0, maximum=100, interactive=True)
                            mech_Si = gr.Number(label=f"Si (wt%)", value=0, minimum=0, maximum=100, interactive=True)
                            mech_Sn = gr.Number(label=f"Sn (wt%)", value=0, minimum=0, maximum=100, interactive=True)
                            mech_V = gr.Number(label=f"V (wt%)", value=0, minimum=0, maximum=100, interactive=True)
                            mech_Zr = gr.Number(label=f"Zr (wt%)", value=0, minimum=0, maximum=100, interactive=True)
                            mech_N = gr.Number(label=f"N (wt%)", value=0, minimum=0, maximum=100, interactive=True)
                            mech_O = gr.Number(label=f"O (wt%)", value=0, minimum=0, maximum=100, interactive=True)
                            mech_C = gr.Number(label=f"C (wt%)", value=0, minimum=0, maximum=100, interactive=True)
                            mech_H = gr.Number(label=f"H (wt%)", value=0, minimum=0, maximum=100, interactive=True)
                            mech_B = gr.Number(label=f"B (wt%)", value=0, minimum=0, maximum=100, interactive=True)
                            mech_Mn = gr.Number(label=f"Mn (wt%)", value=0, minimum=0, maximum=100, interactive=True)
                            mech_Ta = gr.Number(label=f"Ta (wt%)", value=0, minimum=0, maximum=100, interactive=True)
                            mech_Bi = gr.Number(label=f"Bi (wt%)", value=0, minimum=0, maximum=100, interactive=True)
                            mech_Co = gr.Number(label=f"Co (wt%)", value=0, minimum=0, maximum=100, interactive=True)
                        gr.Markdown("### 工艺参数")
                        mech_htt = gr.Number(label="热处理温度 (°C)", value=600, interactive=True)
                        with gr.Row():
                            mech_clear_btn = gr.Button("重置", interactive=True)
                            mech_run_btn = gr.Button("提交", interactive=True)
                    with gr.Column():
                        gr.Markdown("### 力学性能")
                        mech_ys = gr.Number(label="屈服强度 (MPa)", value=None, interactive=False, precision=3)
                        mech_ts = gr.Number(label="抗拉强度 (MPa)", value=None, interactive=False, precision=3)
                        mech_h = gr.Number(label="硬度 (VPN)", value=None, interactive=False, precision=3)
                        mech_hp = gr.Number(label="霍尔佩奇系数 (MPa·m^(1/2))", value=None, interactive=False, precision=3)
            with gr.Tab("批量模式"):
                with gr.Row():
                    with gr.Column():
                        up_files = gr.File(label="上传 CSV 或 XLSX 文件",
                                           file_types=[".csv", ".xlsx"],
                                           file_count="multiple",
                                           type="filepath",
                                           interactive=True)
                        down_files = gr.File(label="下载 CSV 或 XLSX 文件",
                                             file_types=[".csv", ".xlsx"],
                                             file_count="multiple",
                                             type="filepath",
                                             interactive=False)
                    with gr.Column():
                        checkboxes = gr.CheckboxGroup(choices=["热膨胀系数", "密度", "热导率", "电导率", "杨氏模量",
                                                               "体积模量","剪切模量", "泊松比", "比焓", "比热容",
                                                               "屈服强度", "抗拉强度", "硬度", "霍尔佩奇系数"],
                                                      label="选择性能",
                                                      interactive=True)
                        with gr.Row():
                            batch_clear_btn = gr.Button("重置", interactive=True)
                            batch_run_btn = gr.Button("提交", interactive=True)
                            batch_all_btn = gr.Button("全选", interactive=True)
            with gr.Tab("模板下载"):
                gr.File(value=[str(resources.files("tapp.resource").joinpath("tapp_template.csv")),
                               str(resources.files("tapp.resource").joinpath("tapp_template.xlsx"))],
                        label="下载批处理模板文件",
                        file_count="multiple",
                        type="filepath",
                        interactive=False)
        # 定义交互逻辑
        phys_clear_btn.click(
            fn=lambda: [100, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 600,
                        None, None, None, None, None, None, None, None, None, None],
            outputs=[phys_Ti, phys_Al, phys_Cr, phys_Cu, phys_Fe, phys_Mo, phys_Ni, phys_Nb, phys_Si, phys_Sn, phys_V,
                     phys_Zr, phys_N, phys_O, phys_C, phys_H, phys_B, phys_Mn, phys_Ta, phys_Bi, phys_Co, phys_htt,
                     phys_te, phys_d, phys_tc, phys_ec, phys_ym, phys_bm, phys_sm, phys_pr, phys_se, phys_shc]
        )
        phys_run_btn.click(
            fn=get_phys_prop_values,
            inputs=[phys_Ti, phys_Al, phys_Cr, phys_Cu, phys_Fe, phys_Mo, phys_Ni, phys_Nb, phys_Si, phys_Sn, phys_V,
                    phys_Zr, phys_N, phys_O, phys_C, phys_H, phys_B, phys_Mn, phys_Ta, phys_Bi, phys_Co, phys_htt],
            outputs=[phys_te, phys_d, phys_tc, phys_ec, phys_ym, phys_bm, phys_sm, phys_pr, phys_se, phys_shc]
        )
        mech_clear_btn.click(
            fn=lambda: [100, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 600, None, None, None, None],
            outputs=[mech_Ti, mech_Al, mech_Cr, mech_Cu, mech_Fe, mech_Mo, mech_Ni, mech_Nb, mech_Si, mech_Sn, mech_V,
                     mech_Zr, mech_N, mech_O, mech_C, mech_H, mech_B, mech_Mn, mech_Ta, mech_Bi, mech_Co, mech_htt,
                     mech_ys, mech_ts, mech_h, mech_hp]
        )
        mech_run_btn.click(
            fn=get_mech_prop_values,
            inputs=[mech_Ti, mech_Al, mech_Cr, mech_Cu, mech_Fe, mech_Mo, mech_Ni, mech_Nb, mech_Si, mech_Sn, mech_V,
                    mech_Zr, mech_N, mech_O, mech_C, mech_H, mech_B, mech_Mn, mech_Ta, mech_Bi, mech_Co, mech_htt],
            outputs=[mech_ys, mech_ts, mech_h, mech_hp]
        )
        batch_run_btn.click(
            fn=batch_proc,
            inputs=[checkboxes, up_files],
            outputs=down_files
        )
        batch_all_btn.click(
            fn=lambda sel_props: list(
                {"热膨胀系数", "密度", "热导率", "电导率", "杨氏模量", "体积模量", "剪切模量",
                 "泊松比", "比焓","比热容", "屈服强度", "抗拉强度", "硬度", "霍尔佩奇系数"} \
                - set(sel_props)
            ),
            inputs=checkboxes,
            outputs=checkboxes
        )
        batch_clear_btn.click(
            fn=lambda: [None, None, None],
            outputs=[up_files, down_files, checkboxes]
        )
    index.launch()
